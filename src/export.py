"""
Export evaluation results to CSV and Excel.
"""
import csv
import io
from datetime import datetime

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side, GradientFill
    from openpyxl.utils import get_column_letter
    EXCEL_OK = True
except ImportError:
    EXCEL_OK = False


# ── Column definitions ────────────────────────────────────────────────────────

COLUMNS = [
    ("Rank",             "rank"),
    ("Name",             "name"),
    ("Matricule",        "matricule"),
    ("Code (/20)",       "grade_code"),
    ("Execution (/20)",  "grade_execution"),
    ("Docs (/20)",       "grade_documentation"),
    ("Final (/20)",      "grade_final"),
    ("Code Files",       "code_files_count"),
    ("Doc Files",        "doc_files_count"),
    ("Compiled",         "compilation_success"),
    ("Executed",         "execution_success"),
    ("Code URL",         "code_url"),
    ("Doc URL",          "doc_url"),
]


def _row_values(student: dict, rank: int) -> list:
    return [
        rank,
        student.get("name", ""),
        student.get("matricule", ""),
        round(student.get("grade_code", 0), 2),
        round(student.get("grade_execution", 0), 2),
        round(student.get("grade_documentation", 0), 2),
        round(student.get("grade_final", 0), 2),
        student.get("code_files_count", 0),
        student.get("doc_files_count", 0),
        "Yes" if student.get("compilation_success") else "No",
        "Yes" if student.get("execution_success") else "No",
        student.get("code_url", ""),
        student.get("doc_url", ""),
    ]


# ── CSV ───────────────────────────────────────────────────────────────────────

def to_csv(students: list[dict]) -> str:
    """Return CSV string (UTF-8 with BOM for Excel compatibility)."""
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([c[0] for c in COLUMNS])
    sorted_students = sorted(students, key=lambda s: s.get("grade_final", 0), reverse=True)
    for rank, s in enumerate(sorted_students, 1):
        writer.writerow(_row_values(s, rank))
    return "\ufeff" + buf.getvalue()   # BOM


# ── Excel ─────────────────────────────────────────────────────────────────────

def to_excel(students: list[dict], session_name: str = "") -> bytes:
    if not EXCEL_OK:
        raise RuntimeError("openpyxl is not installed")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Evaluation Results"

    # ── Styles ────────────────────────────────────────────────────────────────
    GREEN = "107C10"
    WHITE = "FFFFFF"
    LIGHT = "F1F9F1"
    GRAY  = "F3F2F1"

    header_font   = Font(name="Calibri", bold=True, color=WHITE, size=11)
    header_fill   = PatternFill("solid", fgColor=GREEN)
    header_align  = Alignment(horizontal="center", vertical="center", wrap_text=True)
    center_align  = Alignment(horizontal="center", vertical="center")
    left_align    = Alignment(horizontal="left",   vertical="center")
    thin          = Side(style="thin", color="C8C6C4")
    border        = Border(left=thin, right=thin, top=thin, bottom=thin)

    grade_fills = {
        "excellent": PatternFill("solid", fgColor="107C10"),
        "good":      PatternFill("solid", fgColor="5CA85C"),
        "average":   PatternFill("solid", fgColor="F0A500"),
        "poor":      PatternFill("solid", fgColor="C50F1F"),
    }
    grade_fonts = {
        "excellent": Font(bold=True, color=WHITE),
        "good":      Font(bold=True, color=WHITE),
        "average":   Font(bold=True, color=WHITE),
        "poor":      Font(bold=True, color=WHITE),
    }

    def grade_tier(v):
        if v >= 16:  return "excellent"
        if v >= 12:  return "good"
        if v >= 10:  return "average"
        return "poor"

    # ── Title row ────────────────────────────────────────────────────────────
    ws.merge_cells("A1:M1")
    title_cell = ws["A1"]
    title_cell.value = f"GitHub Evaluator — {session_name or 'Results'} — {datetime.now().strftime('%Y-%m-%d')}"
    title_cell.font = Font(name="Calibri", bold=True, size=14, color="323130")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    title_cell.fill = PatternFill("solid", fgColor=LIGHT)
    ws.row_dimensions[1].height = 28

    # ── Header row ───────────────────────────────────────────────────────────
    for col, (label, _) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=2, column=col, value=label)
        cell.font   = header_font
        cell.fill   = header_fill
        cell.alignment = header_align
        cell.border = border
    ws.row_dimensions[2].height = 32

    # ── Data rows ────────────────────────────────────────────────────────────
    sorted_students = sorted(students, key=lambda s: s.get("grade_final", 0), reverse=True)
    for row_i, s in enumerate(sorted_students, 3):
        rank = row_i - 2
        values = _row_values(s, rank)
        fill_row = PatternFill("solid", fgColor=(LIGHT if rank % 2 == 0 else WHITE))

        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row_i, column=col, value=val)
            cell.border  = border
            cell.fill    = fill_row
            cell.alignment = center_align if col != 2 else left_align

        # Colour the Final grade cell
        final_cell = ws.cell(row=row_i, column=7)
        tier = grade_tier(s.get("grade_final", 0))
        final_cell.fill = grade_fills[tier]
        final_cell.font = grade_fonts[tier]

        ws.row_dimensions[row_i].height = 22

    # ── Statistics ───────────────────────────────────────────────────────────
    if students:
        stats_row = len(students) + 4
        ws.cell(stats_row, 1, "Statistics").font = Font(bold=True, size=11)
        finals = [s.get("grade_final", 0) for s in students]
        stats = [
            ("Average", round(sum(finals) / len(finals), 2)),
            ("Maximum", round(max(finals), 2)),
            ("Minimum", round(min(finals), 2)),
            ("Pass rate", f"{round(sum(1 for f in finals if f >= 10) / len(finals) * 100, 1)}%"),
        ]
        for i, (label, val) in enumerate(stats):
            ws.cell(stats_row + 1 + i, 1, label).font = Font(bold=True)
            ws.cell(stats_row + 1 + i, 2, val)

    # ── Column widths ────────────────────────────────────────────────────────
    widths = [6, 28, 16, 12, 15, 12, 12, 11, 10, 10, 10, 40, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ── Freeze header ────────────────────────────────────────────────────────
    ws.freeze_panes = "A3"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
